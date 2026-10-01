"""Reviewed offline workbook generator/reference. Executed only in the recipe container."""
import base64
import io
import json
import random
import sys
import zipfile
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

HEADERS = ['ID', 'Date', 'Quantity', 'Price', 'Amount']
TASK = ('Clean Data rows: preserve first occurrence of each ID; discard rows with invalid numeric '
        'quantity/price; missing numbers become zero. Preserve order and ID text. Normalize dates '
        'to YYYY-MM-DD text. Set Amount to a relative Quantity*Price formula on every retained row. '
        'Write Summary!A1="Total" and Summary!B1=SUM(Data!E2:E<last-row>). Preserve other sheets '
        'and their values. Keep all Data header cells bold. Save an XLSX with recalculated formula '
        'caches; return its base64 bytes as workbook_b64. No macros or external links.')
VARIATIONS = ['mixed_types_dates', 'missing_invalid', 'duplicates', 'multiple_sheets',
              'formulas', 'structure_formatting', 'unicode_sparse']


def encode(book):
    book.properties.created = book.properties.modified = datetime(2000, 1, 1)
    stream = io.BytesIO()
    book.save(stream)
    return canonical(stream.getvalue())


def canonical(raw):
    target = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw)) as source, zipfile.ZipFile(target, 'w') as dest:
        for name in sorted(source.namelist()):
            value = source.read(name)
            if name == 'docProps/core.xml':
                from xml.etree import ElementTree as ET
                root = ET.fromstring(value)
                for field in ('created', 'modified'):
                    root.find('{http://purl.org/dc/terms/}' + field).text = '2000-01-01T00:00:00Z'
                value = ET.tostring(root)
            info = zipfile.ZipInfo(name, (2000, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            dest.writestr(info, value)
    return base64.b64encode(target.getvalue()).decode()


def render(rows, extras):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Data'
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for name, value in extras.items():
        book.create_sheet(name)['A1'] = value
    return encode(book)


def clean(rows):
    result, seen = [], set()
    for identity, date, quantity, price, *_ in rows:
        if identity is None or identity in seen:
            continue
        seen.add(identity)
        try:
            quantity = int(quantity or 0)
            price = int(price or 0)
            date = datetime.strptime(str(date)[:10], '%Y-%m-%d').strftime('%Y-%m-%d')
        except (TypeError, ValueError):
            continue
        result.append([str(identity), date, quantity, price, quantity * price])
    return result


def generate(payload):
    rng = random.Random(payload['seed'])
    cases = []
    for stage in payload['stages']:
        for index in range(stage['count']):
            variation = VARIATIONS[index % len(VARIATIONS)]
            rows = [[f'{stage["split"]}-{rng.getrandbits(48):012x}-{i}', f'2026-09-{i+1:02}',
                     rng.randrange(1, 20), rng.randrange(1, 100), None] for i in range(2 + index % 14)]
            extras = {'Notes': 'Preserve this note'} if variation in {'multiple_sheets', 'structure_formatting'} else {}
            if variation == 'mixed_types_dates':
                for row in rows[::2]:
                    row[2] = str(row[2])
                    row[1] = datetime.strptime(row[1], '%Y-%m-%d')
            elif variation == 'missing_invalid':
                rows[0][2] = None
                rows[-1][3] = 'invalid'
            elif variation == 'duplicates':
                rows.insert(1, list(rows[0]))
                rows[1][3] = 999
            elif variation == 'formulas':
                for row in rows:
                    row[4] = '=1+1'
            elif variation == 'unicode_sparse':
                rows[0][0] = '東京\u2013Zoë-' + rows[0][0]
                rows.insert(1, [None] * 5)
            truth = {'rows': clean(rows), 'extras': extras}
            value = {'workbook_b64': render(rows, extras), 'task': TASK,
                     'variation': variation, 'row_bound': len(rows)}
            cases.append({'input': value, 'output': truth, 'variation': variation,
                          'template_family': 'workbook:' + variation})
    return {'cases': cases}


def reference(value):
    # Read the actual rendered input; generator truth is never an input here.
    book = load_workbook(io.BytesIO(base64.b64decode(value['workbook_b64'])), data_only=False)
    rows = clean(list(book['Data'].values)[1:])
    sheet = book['Data']
    sheet.delete_rows(2, sheet.max_row)
    for index, row in enumerate(rows, 2):
        sheet.append([*row[:4], f'=C{index}*D{index}'])
    if 'Summary' not in book:
        book.create_sheet('Summary')
    book['Summary']['A1'] = 'Total'
    book['Summary']['B1'] = f'=SUM(Data!E2:E{len(rows)+1})'
    # openpyxl writes formulas without caches. Populate this deterministic
    # reference's caches; the separate grader recomputes from actual cells.
    from xml.etree import ElementTree as ET
    raw = base64.b64decode(encode(book))
    target = io.BytesIO()
    ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    with zipfile.ZipFile(io.BytesIO(raw)) as source, zipfile.ZipFile(target, 'w', zipfile.ZIP_DEFLATED) as dest:
        for name in source.namelist():
            data = source.read(name)
            if name.startswith('xl/worksheets/'):
                root = ET.fromstring(data)
                for cell in root.iter(ns+'c'):
                    formula = cell.find(ns+'f')
                    if formula is None:
                        continue
                    cached = cell.find(ns+'v')
                    if cached is None:
                        cached = ET.SubElement(cell, ns+'v')
                    cached.text = str(sum(row[4] for row in rows) if formula.text.startswith('SUM(')
                                      else rows[int(cell.attrib['r'][1:])-2][4])
                data = ET.tostring(root)
            dest.writestr(name, data)
    return {'workbook_b64': canonical(target.getvalue())}


if __name__ == '__main__':
    payload = json.load(sys.stdin)
    print(json.dumps(generate(payload) if sys.argv[1] == 'generate'
                     else {'outputs': [reference(value) for value in payload['inputs']]}))
